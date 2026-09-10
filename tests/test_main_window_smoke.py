from os.path import normpath
from unittest.mock import patch

from PyQt6.QtCore import QMimeData, QUrl
from PyQt6.QtGui import QColor, QDesktopServices
from PyQt6.QtWidgets import QCheckBox, QDialog, QFileDialog, QMessageBox, QTableWidgetItem

from main_window import FILE_PATH_ROLE, FileScanner
from scanner import MatchMode, Occurrence


class FakeDropEvent:
    def __init__(self, mime_data):
        self._mime_data = mime_data

    def mimeData(self):
        return self._mime_data

    def acceptProposedAction(self):
        pass


def test_window_launches(qtbot):
    window = FileScanner()
    qtbot.addWidget(window)

    assert window.windowTitle() == 'File Scanner'


def test_window_icon_loads_from_assets(qtbot):
    window = FileScanner()
    qtbot.addWidget(window)

    assert not window.windowIcon().isNull()


def test_scan_button_enabled_after_keyword_then_file(qtbot, tmp_path, monkeypatch):
    a_file = str(tmp_path / 'a.txt')
    (tmp_path / 'a.txt').write_text('placeholder')
    monkeypatch.setattr(QFileDialog, 'getOpenFileNames', staticmethod(lambda *a, **k: ([a_file], '')))

    window = FileScanner()
    qtbot.addWidget(window)
    assert not window.scan_files.isEnabled()

    window.new_keyword_text.setText('apple')
    window.add_keyword_clicked()
    assert not window.scan_files.isEnabled()

    window.add_files_clicked()
    assert window.scan_files.isEnabled()


def test_scan_button_enabled_after_file_then_keyword(qtbot, tmp_path, monkeypatch):
    a_file = str(tmp_path / 'a.txt')
    (tmp_path / 'a.txt').write_text('placeholder')
    monkeypatch.setattr(QFileDialog, 'getOpenFileNames', staticmethod(lambda *a, **k: ([a_file], '')))

    window = FileScanner()
    qtbot.addWidget(window)
    assert not window.scan_files.isEnabled()

    window.add_files_clicked()
    assert not window.scan_files.isEnabled()

    window.new_keyword_text.setText('apple')
    window.add_keyword_clicked()
    assert window.scan_files.isEnabled()


def test_scan_populates_results_and_reports_per_file_errors(qtbot, tmp_path, monkeypatch):
    monkeypatch.setattr(QMessageBox, 'warning', staticmethod(lambda *a, **k: None))

    good_file = tmp_path / 'good.txt'
    good_file.write_text('this file mentions apple')
    bad_file = tmp_path / 'bad.docx'
    bad_file.write_bytes(b'not a real docx')

    window = FileScanner()
    qtbot.addWidget(window)
    window.keyword_list.addItem('apple')
    window.update_file_headers()
    window._add_files([str(good_file), str(bad_file)])
    window.scan_files.setDisabled(False)

    window.scan_files_clicked()
    qtbot.waitUntil(lambda: not window.scan_worker.isRunning(), timeout=15000)
    qtbot.wait(50)

    assert window.files.item(0, 1).text() == 'true'
    assert window.files.item(1, 1).text() == 'ERROR'
    assert len(window.scan_errors) == 1
    assert window.export_results.isEnabled()
    assert window.scan_files.isEnabled()


def test_scan_delivers_results_in_batches_and_populates_every_row(qtbot, tmp_path, monkeypatch):
    # Force a tiny batch so a handful of files still spans several batches.
    monkeypatch.setattr('main_window.RESULTS_BATCH_FRACTION', 0.25)

    files = []
    has_apple = {}
    for i in range(8):
        path = tmp_path / f'doc{i}.txt'
        path.write_text('apple' if i % 2 else 'orange')
        files.append(str(path))
        has_apple[str(path)] = bool(i % 2)

    window = FileScanner()
    qtbot.addWidget(window)
    window.keyword_list.addItem('apple')
    window.update_file_headers()
    window._add_files(files)
    window.scan_files.setDisabled(False)

    batch_sizes = []
    original = window.handle_scan_batch

    def spy(results):
        batch_sizes.append(len(results))
        original(results)

    monkeypatch.setattr(window, 'handle_scan_batch', spy)

    window.scan_files_clicked()
    qtbot.waitUntil(lambda: not window.scan_worker.isRunning(), timeout=15000)
    qtbot.wait(50)

    assert sum(batch_sizes) == len(files)
    assert len(batch_sizes) > 1  # delivered in batches, not one giant emission
    assert max(batch_sizes) <= 2  # batch_size == int(8 * 0.25)
    for row in range(window.files.rowCount()):
        file = window.files.item(row, 0).data(FILE_PATH_ROLE)
        expected = 'true' if has_apple[file] else 'false'
        assert window.files.item(row, 1).text() == expected


def test_remove_files_clicked_removes_selected_rows(qtbot, tmp_path):
    a_file = str(tmp_path / 'a.txt')
    b_file = str(tmp_path / 'b.txt')
    (tmp_path / 'a.txt').write_text('placeholder')
    (tmp_path / 'b.txt').write_text('placeholder')

    window = FileScanner()
    qtbot.addWidget(window)
    window._add_files([a_file, b_file])
    assert not window.remove_files.isEnabled()

    window.files.selectRow(0)
    assert window.remove_files.isEnabled()

    window.remove_files_clicked()

    assert window.file_names == [b_file]
    assert window.files.rowCount() == 1
    assert not window.remove_files.isEnabled()


def test_drag_and_drop_adds_files(qtbot, tmp_path):
    a_file = tmp_path / 'a.txt'
    a_file.write_text('placeholder')

    window = FileScanner()
    qtbot.addWidget(window)

    mime_data = QMimeData()
    mime_data.setUrls([QUrl.fromLocalFile(str(a_file))])
    window.dropEvent(FakeDropEvent(mime_data))

    assert [normpath(f) for f in window.file_names] == [normpath(str(a_file))]
    assert window.files.rowCount() == 1
    assert window.scan_files.isEnabled() is False  # no keywords added yet


def test_cancel_mid_scan_stops_worker_and_resets_buttons(qtbot, tmp_path):
    files = []
    for i in range(20):
        path = tmp_path / f'f{i}.txt'
        path.write_text('hello world ' * 1000)
        files.append(str(path))

    window = FileScanner()
    qtbot.addWidget(window)
    window.keyword_list.addItem('hello')
    window.update_file_headers()
    window._add_files(files)
    window.scan_files.setDisabled(False)

    window.scan_files_clicked()
    qtbot.wait(10)
    window.progress_dialog.cancel()
    qtbot.waitUntil(lambda: not window.scan_worker.isRunning(), timeout=15000)

    assert window.scan_files.isEnabled()
    assert window.add_files.isEnabled()


def test_whole_word_mode_used_during_scan(qtbot, tmp_path):
    path = tmp_path / 'notes.txt'
    path.write_text('concatenate this')

    window = FileScanner()
    qtbot.addWidget(window)
    window.keyword_list.addItem('cat')
    window.update_file_headers()
    window.match_mode_combo.setCurrentIndex(window.match_mode_combo.findData(MatchMode.WHOLE_WORD))
    window._add_files([str(path)])
    window.scan_files.setDisabled(False)

    window.scan_files_clicked()
    qtbot.waitUntil(lambda: not window.scan_worker.isRunning(), timeout=15000)
    qtbot.wait(50)

    assert window.files.item(0, 1).text() == 'false'


def test_scan_files_clicked_rejects_invalid_regex_before_starting(qtbot, monkeypatch):
    monkeypatch.setattr(QMessageBox, 'critical', staticmethod(lambda *a, **k: None))

    window = FileScanner()
    qtbot.addWidget(window)
    window.keyword_list.addItem('(')
    window.update_file_headers()
    window.match_mode_combo.setCurrentIndex(window.match_mode_combo.findData(MatchMode.REGEX))
    window.file_names = ['dummy.txt']
    window.scan_files.setDisabled(False)

    window.scan_files_clicked()

    assert window.scan_worker is None


def test_regex_mode_flags_invalid_text_as_typed(qtbot):
    window = FileScanner()
    qtbot.addWidget(window)
    window.match_mode_combo.setCurrentIndex(window.match_mode_combo.findData(MatchMode.REGEX))

    window.new_keyword_text.setText('(')
    window._update_keyword_input_validity()

    assert not window.add_keyword.isEnabled()
    style = window.new_keyword_text.styleSheet()
    # Both colors must be set explicitly - a background-only stylesheet left
    # the text color to the OS theme, unreadable in dark mode.
    assert window.options.display.invalid_regex_background.name() in style
    assert window.options.display.invalid_regex_text.name() in style

    window.new_keyword_text.setText('valid')
    window._update_keyword_input_validity()

    assert window.add_keyword.isEnabled()
    assert window.new_keyword_text.styleSheet() == ''


def test_substring_mode_does_not_flag_regex_metacharacters(qtbot):
    window = FileScanner()
    qtbot.addWidget(window)

    window.new_keyword_text.setText('(')
    window._update_keyword_input_validity()

    assert window.add_keyword.isEnabled()
    assert window.new_keyword_text.styleSheet() == ''


def test_switching_to_regex_mode_highlights_existing_invalid_keyword_and_disables_scan(qtbot, tmp_path):
    a_file = str(tmp_path / 'a.txt')
    (tmp_path / 'a.txt').write_text('placeholder')

    window = FileScanner()
    qtbot.addWidget(window)
    window.keyword_list.addItem('(')
    window.update_file_headers()
    window._add_files([a_file])
    assert window.scan_files.isEnabled()  # substring mode - '(' is just a literal character

    window.match_mode_combo.setCurrentIndex(window.match_mode_combo.findData(MatchMode.REGEX))

    assert not window.scan_files.isEnabled()
    item = window.keyword_list.item(0)
    assert item.background().color() == window.options.display.invalid_regex_background
    assert item.foreground().color() == window.options.display.invalid_regex_text

    window.match_mode_combo.setCurrentIndex(window.match_mode_combo.findData(MatchMode.SUBSTRING))

    assert window.scan_files.isEnabled()
    assert item.background().color() != window.options.display.invalid_regex_background


def test_keyword_list_is_not_editable(qtbot):
    window = FileScanner()
    qtbot.addWidget(window)

    assert window.keyword_list.editTriggers() == window.keyword_list.EditTrigger.NoEditTriggers


def test_export_excel_writes_workbook_with_colors(qtbot, tmp_path, monkeypatch):
    from openpyxl import load_workbook

    xlsx_path = str(tmp_path / 'out.xlsx')
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', staticmethod(lambda *a, **k: (xlsx_path, '')))

    window = FileScanner()
    qtbot.addWidget(window)
    window.keyword_list.addItem('apple')
    window.update_file_headers()
    window.files.setRowCount(1)
    window.files.setItem(0, 0, QTableWidgetItem('good.txt'))
    found_item = QTableWidgetItem('true')
    found_item.setBackground(window.options.display.found_color)
    window.files.setItem(0, 1, found_item)

    window.export_excel_clicked()

    workbook = load_workbook(xlsx_path)
    sheet = workbook.active
    assert [cell.value for cell in sheet[1]] == ['File', 'apple']
    assert [cell.value for cell in sheet[2]] == ['good.txt', 'true']
    hex_color = window.options.display.found_color.name(QColor.NameFormat.HexRgb).lstrip('#').upper()
    assert hex_color in sheet.cell(row=2, column=2).fill.fgColor.rgb.upper()
    # the filename column never has an explicit background set - it must not
    # get colored just because an unset QBrush reports an opaque black color
    assert sheet.cell(row=2, column=1).fill.fill_type is None


def test_export_save_dialog_appends_missing_extension(qtbot, tmp_path, monkeypatch):
    from openpyxl import load_workbook

    typed_path = str(tmp_path / 'results')  # no extension typed
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', staticmethod(lambda *a, **k: (typed_path, '')))

    window = FileScanner()
    qtbot.addWidget(window)
    window.export_excel_clicked()

    load_workbook(typed_path + '.xlsx')  # raises if the file wasn't created with the extension


def test_export_excel_strips_illegal_xml_characters_from_snippets(qtbot, tmp_path, monkeypatch):
    from openpyxl import load_workbook

    xlsx_path = str(tmp_path / 'out.xlsx')
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', staticmethod(lambda *a, **k: (xlsx_path, '')))

    a_file = str(tmp_path / 'a.txt')
    window = FileScanner()
    qtbot.addWidget(window)
    window.keyword_list.addItem('apple')
    window.update_file_headers()
    window._add_files([a_file])
    window.files.setItem(0, 1, QTableWidgetItem('true'))
    # A control char and the noncharacter U+FFFE - both illegal in XML 1.0.
    # pypdfium2 can produce these as decoding garbage from a real PDF; this
    # simulates that without needing to craft one.
    illegal_snippet = 'before' + chr(0x00) + 'after' + chr(0xFFFE) + 'end'
    window.file_occurrences[a_file] = [[Occurrence(page=1, snippet=illegal_snippet)]]

    def fake_exec(self):
        for checkbox in self.findChildren(QCheckBox):
            checkbox.setChecked(True)
        return QDialog.DialogCode.Accepted

    with patch.object(QDialog, 'exec', fake_exec):
        window.export_excel_clicked()  # must not raise, and must produce a loadable file

    workbook = load_workbook(xlsx_path)  # raises if the file is corrupt
    sheet = workbook.active
    snippet_cell = sheet.cell(row=2, column=4).value  # File, apple, apple (page), apple (snippet)
    assert snippet_cell == 'beforeafterend'


def test_ask_export_extra_columns_skips_dialog_when_no_occurrences(qtbot, monkeypatch):
    window = FileScanner()
    qtbot.addWidget(window)
    monkeypatch.setattr(QDialog, 'exec', lambda self: (_ for _ in ()).throw(AssertionError('dialog should not show')))

    assert window._ask_export_extra_columns() == (False, False)


def test_ask_export_extra_columns_returns_checked_state(qtbot, tmp_path):
    a_file = str(tmp_path / 'a.txt')

    window = FileScanner()
    qtbot.addWidget(window)
    window.file_occurrences[a_file] = [[Occurrence(page=1, snippet='hello')]]

    def fake_exec(self):
        for checkbox in self.findChildren(QCheckBox):
            checkbox.setChecked(True)
        return QDialog.DialogCode.Accepted

    with patch.object(QDialog, 'exec', fake_exec):
        assert window._ask_export_extra_columns() == (True, True)

    with patch.object(QDialog, 'exec', lambda self: QDialog.DialogCode.Rejected):
        assert window._ask_export_extra_columns() is None


def test_export_rows_include_page_and_snippet_columns(qtbot, tmp_path):
    path = tmp_path / 'notes.txt'
    path.write_text('apple pie')

    window = FileScanner()
    qtbot.addWidget(window)
    window.keyword_list.addItem('apple')
    window.update_file_headers()
    window._add_files([str(path)])
    window.file_occurrences[str(path)] = [[Occurrence(page=None, snippet='apple pie')]]
    window.files.setItem(0, 1, QTableWidgetItem('true'))

    rows = list(window._export_rows(include_pages=True, include_snippets=True))

    assert rows[0] == ['File', 'apple', 'apple (page)', 'apple (snippet)']
    assert rows[1] == ['notes.txt', 'true', '', 'apple pie']


def test_window_geometry_persists_across_instances(qtbot):
    window = FileScanner()
    qtbot.addWidget(window)
    window.resize(800, 600)
    window.close()

    reopened = FileScanner()
    qtbot.addWidget(reopened)

    # offscreen QPA rounds frame geometry by a couple of pixels, so allow slack
    assert abs(reopened.size().width() - 800) <= 5
    assert abs(reopened.size().height() - 600) <= 5


def test_display_options_persist_across_instances(qtbot):
    window = FileScanner()
    qtbot.addWidget(window)
    window.options.display.found_color = QColor('#123456')
    window.options.display.error_text = 'BROKEN'
    window.close()

    reopened = FileScanner()
    qtbot.addWidget(reopened)

    assert reopened.options.display.found_color == QColor('#123456')
    assert reopened.options.display.error_text == 'BROKEN'
    # the dialog's own widgets must reflect the loaded values too
    assert reopened.options.error_text_input.text() == 'BROKEN'


def test_custom_error_color_and_text_used_in_scan_results(qtbot, tmp_path, monkeypatch):
    monkeypatch.setattr(QMessageBox, 'warning', staticmethod(lambda *a, **k: None))
    bad_file = tmp_path / 'bad.docx'
    bad_file.write_bytes(b'not a real docx')

    window = FileScanner()
    qtbot.addWidget(window)
    window.options.display.error_text = 'BROKEN'
    window.options.display.error_color = QColor('purple')
    window.keyword_list.addItem('apple')
    window.update_file_headers()
    window._add_files([str(bad_file)])
    window.scan_files.setDisabled(False)

    window.scan_files_clicked()
    qtbot.waitUntil(lambda: not window.scan_worker.isRunning(), timeout=15000)
    qtbot.wait(50)

    item = window.files.item(0, 1)
    assert item.text() == 'BROKEN'
    assert item.background().color() == QColor('purple')


def test_dialogs_open_in_last_used_directory(qtbot, tmp_path, monkeypatch):
    used_dirs = []
    a_file = str(tmp_path / 'a.txt')
    (tmp_path / 'a.txt').write_text('placeholder')

    def fake_get_open_file_names(_parent, _title, directory, _filter):
        used_dirs.append(directory)
        return [a_file], ''

    monkeypatch.setattr(QFileDialog, 'getOpenFileNames', staticmethod(fake_get_open_file_names))

    window = FileScanner()
    qtbot.addWidget(window)
    window.add_files_clicked()
    window.add_files_clicked()

    assert used_dirs[0] == '.'
    assert used_dirs[1] == str(tmp_path)


def test_found_cell_tooltip_shows_snippet(qtbot, tmp_path):
    path = tmp_path / 'notes.txt'
    path.write_text('this file mentions apple pie')

    window = FileScanner()
    qtbot.addWidget(window)
    window.keyword_list.addItem('apple')
    window.update_file_headers()
    window._add_files([str(path)])
    window.scan_files.setDisabled(False)

    window.scan_files_clicked()
    qtbot.waitUntil(lambda: not window.scan_worker.isRunning(), timeout=15000)
    qtbot.wait(50)

    tooltip = window.files.item(0, 1).toolTip()
    assert 'apple' in tooltip
    assert 'p.' not in tooltip  # .txt has no page number, so no "p. N:" prefix


def test_sorting_by_result_column_keeps_correct_file_per_row(qtbot, tmp_path):
    file_a = tmp_path / 'a.txt'
    file_a.write_text('no keyword here')
    file_b = tmp_path / 'b.txt'
    file_b.write_text('apple is here')

    window = FileScanner()
    qtbot.addWidget(window)
    window.keyword_list.addItem('apple')
    window.update_file_headers()
    window._add_files([str(file_a), str(file_b)])
    window.scan_files.setDisabled(False)

    window.scan_files_clicked()
    qtbot.waitUntil(lambda: not window.scan_worker.isRunning(), timeout=15000)
    qtbot.wait(50)

    window.files.sortItems(1)  # sort by the result column, not just filename

    for row in range(window.files.rowCount()):
        file = window.files.item(row, 0).data(FILE_PATH_ROLE)
        expected_text = 'true' if file == str(file_b) else 'false'
        assert window.files.item(row, 1).text() == expected_text

    # remove_files_clicked must still target the right file after the sort
    for row in range(window.files.rowCount()):
        if window.files.item(row, 0).data(FILE_PATH_ROLE) == str(file_a):
            window.files.selectRow(row)
            break
    window.remove_files_clicked()

    assert window.file_names == [str(file_b)]


def test_filter_hides_non_matching_rows(qtbot, tmp_path):
    file_a = tmp_path / 'alpha.txt'
    file_a.write_text('x')
    file_b = tmp_path / 'beta.txt'
    file_b.write_text('x')

    window = FileScanner()
    qtbot.addWidget(window)
    window._add_files([str(file_a), str(file_b)])

    window.filter_text.setText('alpha')

    for row in range(window.files.rowCount()):
        is_alpha = window.files.item(row, 0).text() == 'alpha.txt'
        assert window.files.isRowHidden(row) == (not is_alpha)

    window.filter_text.setText('')

    assert not window.files.isRowHidden(0)
    assert not window.files.isRowHidden(1)


def test_double_click_opens_file(qtbot, tmp_path, monkeypatch):
    a_file = tmp_path / 'a.txt'
    a_file.write_text('placeholder')
    opened = []
    monkeypatch.setattr(QDesktopServices, 'openUrl', staticmethod(lambda url: opened.append(url.toLocalFile())))

    window = FileScanner()
    qtbot.addWidget(window)
    window._add_files([str(a_file)])

    window._open_file_at_row(0, 0)

    assert [normpath(path) for path in opened] == [normpath(str(a_file))]


def test_double_click_missing_file_warns_instead_of_opening(qtbot, tmp_path, monkeypatch):
    a_file = tmp_path / 'a.txt'
    a_file.write_text('placeholder')
    warned = []
    monkeypatch.setattr(QMessageBox, 'warning', staticmethod(lambda *a, **k: warned.append(True)))
    opened = []
    monkeypatch.setattr(QDesktopServices, 'openUrl', staticmethod(lambda url: opened.append(url)))

    window = FileScanner()
    qtbot.addWidget(window)
    window._add_files([str(a_file)])
    a_file.unlink()

    window._open_file_at_row(0, 0)

    assert warned == [True]
    assert opened == []
