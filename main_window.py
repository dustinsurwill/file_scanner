import re
import sys
from _csv import writer
from functools import partial
from multiprocessing import cpu_count, get_context
from os.path import basename, dirname, isfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from PyQt6.QtCore import QSettings, Qt, QThread, QUrl, pyqtSignal
from PyQt6.QtGui import QBrush, QColor, QDesktopServices, QIcon
from PyQt6.QtWidgets import (
    QApplication,
    QComboBox,
    QFileDialog,
    QGridLayout,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QListWidget,
    QMainWindow,
    QMessageBox,
    QProgressDialog,
    QPushButton,
    QStyleFactory,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from options_dialog import Options
from scanner import MatchMode, ScanResult, scan_files_process

FILE_PATH_ROLE = Qt.ItemDataRole.UserRole

# Recycle each worker after this many files so any per-file memory that a
# parsing library doesn't fully release (observed with pdfminer.six on large
# PDF batches) can't accumulate across the life of a long-running scan.
MAX_TASKS_PER_CHILD = 50
# Resolves correctly both from source (relative to this file) and in a
# Nuitka onefile build, where --include-data-files places it at this same
# relative path inside the runtime extraction directory.
ICON_PATH = str(Path(__file__).resolve().parent / 'assets' / 'icon.png')


def _build_tooltip(occurrences, truncated):
    lines = [f'p. {occ.page}: {occ.snippet}' if occ.page is not None else occ.snippet for occ in occurrences]
    if truncated:
        lines.append(f'…and more matches not shown (showing first {len(occurrences)})')
    return '\n'.join(lines)


class ScanWorker(QThread):
    result_ready = pyqtSignal(object)
    finished_scanning = pyqtSignal()

    def __init__(self, keywords, file_names, mode=MatchMode.SUBSTRING, parent=None):
        super().__init__(parent)
        self.keywords = keywords
        self.file_names = file_names
        self.mode = mode
        self._pool = None

    def run(self):
        worker_count = min(cpu_count(), len(self.file_names))
        scan = partial(scan_files_process, self.keywords, mode=self.mode)
        # 'spawn' (not the platform-default 'fork' on Linux) avoids worker
        # processes inheriting this QThread's/QApplication's state, which
        # otherwise deadlocks pool shutdown when a Pool is created from a
        # non-main thread in a Qt app.
        context = get_context('spawn')
        with context.Pool(processes=worker_count, maxtasksperchild=MAX_TASKS_PER_CHILD) as pool:
            self._pool = pool
            try:
                for result in pool.imap_unordered(scan, self.file_names):
                    self.result_ready.emit(result)
            except Exception:  # noqa: BLE001, S110 - a cancel-triggered pool.terminate() surfaces here
                pass
        self.finished_scanning.emit()

    def cancel(self):
        if self._pool is not None:
            self._pool.terminate()


class FileScanner(QMainWindow):
    def __init__(self, parent=None):
        super().__init__(parent)
        center = QWidget()
        main_layout = QGridLayout(center)
        main_layout.addLayout(self.create_keyword_area(), 0, 0)
        main_layout.setColumnMinimumWidth(0, 350)
        self.file_headers = ['files']
        main_layout.addLayout(self.create_files_area(), 0, 1)
        main_layout.setColumnMinimumWidth(1, 650)
        main_layout.setColumnStretch(1, 1)
        main_layout.setRowMinimumHeight(0, 500)
        self.setCentralWidget(center)
        self.setWindowTitle('File Scanner')
        if isfile(ICON_PATH):
            self.setWindowIcon(QIcon(ICON_PATH))
        self.setAcceptDrops(True)
        if sys.platform == 'win32':
            # 'windows11' (Qt 6.7+) supports the Windows dark/light color
            # scheme. Fall back to 'fusion' (also dark-mode aware) rather
            # than 'windowsvista' (Qt's old default, which ignores dark mode
            # entirely) in case 'windows11' isn't available for some reason.
            QApplication.setStyle(QStyleFactory.create('windows11') or QStyleFactory.create('fusion'))
        self.file_names = []
        self.file_items = {}
        self.scan_errors = []
        self.scan_worker = None
        self.progress_dialog = None
        self.settings = QSettings('file-scanner', 'FileScanner')
        self.options = Options(self)
        self._load_display_options()
        geometry = self.settings.value('window_geometry')
        if geometry is not None:
            self.restoreGeometry(geometry)

    # Display-option fields persisted via QSettings, alongside window
    # geometry and the last-used directory. Colors round-trip as hex strings.
    _DISPLAY_COLOR_KEYS = (
        'found_color',
        'missing_color',
        'invalid_regex_background',
        'invalid_regex_text',
        'error_color',
    )
    _DISPLAY_TEXT_KEYS = ('found_text', 'missing_text', 'error_text')

    def _load_display_options(self):
        display = self.options.display
        for key in self._DISPLAY_COLOR_KEYS:
            value = self.settings.value(f'display/{key}')
            if value:
                setattr(display, key, QColor(value))
        for key in self._DISPLAY_TEXT_KEYS:
            value = self.settings.value(f'display/{key}')
            if value:
                setattr(display, key, value)
        self.options.refresh_widgets()

    def _save_display_options(self):
        display = self.options.display
        for key in self._DISPLAY_COLOR_KEYS:
            self.settings.setValue(f'display/{key}', getattr(display, key).name())
        for key in self._DISPLAY_TEXT_KEYS:
            self.settings.setValue(f'display/{key}', getattr(display, key))

    def closeEvent(self, event):
        self.settings.setValue('window_geometry', self.saveGeometry())
        self._save_display_options()
        super().closeEvent(event)

    def _last_dir(self):
        return self.settings.value('last_dir', '.')

    def _remember_dir(self, path):
        self.settings.setValue('last_dir', dirname(path))

    def create_files_area(self):
        vertical = QVBoxLayout()
        top_buttons = QHBoxLayout()
        self.add_files = QPushButton('Add Files')
        self.add_files.clicked.connect(self.add_files_clicked)
        top_buttons.addWidget(self.add_files)
        self.remove_files = QPushButton('Remove Files')
        self.remove_files.setDisabled(True)
        self.remove_files.clicked.connect(self.remove_files_clicked)
        top_buttons.addWidget(self.remove_files)
        self.export_results = QPushButton('Export to CSV')
        self.export_results.setDisabled(True)
        self.export_results.clicked.connect(self.export_results_clicked)
        top_buttons.addWidget(self.export_results)
        self.export_excel = QPushButton('Export to Excel')
        self.export_excel.setDisabled(True)
        self.export_excel.clicked.connect(self.export_excel_clicked)
        top_buttons.addWidget(self.export_excel)
        vertical.addLayout(top_buttons)
        filter_row = QHBoxLayout()
        filter_row.addWidget(QLabel('Filter:'))
        self.filter_text = QLineEdit()
        self.filter_text.setPlaceholderText('Type to filter rows...')
        self.filter_text.textChanged.connect(self._apply_filter)
        filter_row.addWidget(self.filter_text)
        vertical.addLayout(filter_row)
        self.files = QTableWidget()
        self.files.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.files.setAlternatingRowColors(True)
        self.files.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.files.setSortingEnabled(True)
        self.files.itemSelectionChanged.connect(
            lambda: self.remove_files.setDisabled(not self.files.selectedIndexes())
        )
        self.files.cellDoubleClicked.connect(self._open_file_at_row)
        self.files.horizontalHeader().sortIndicatorChanged.connect(self._apply_filter)
        self.update_file_headers()
        vertical.addWidget(self.files)
        buttons = QHBoxLayout()
        self.save_keywords = QPushButton('Save Keywords')
        self.save_keywords.clicked.connect(self.save_keywords_clicked)
        self.save_keywords.setDisabled(True)
        buttons.addWidget(self.save_keywords)
        load_keywords = QPushButton('Load Keywords')
        load_keywords.clicked.connect(self.load_keywords_clicked)
        buttons.addWidget(load_keywords)
        options = QPushButton('Options')
        options.clicked.connect(lambda: self.options.show())
        buttons.addWidget(options)
        self.scan_files = QPushButton('Scan Files')
        self.scan_files.setDisabled(True)
        self.scan_files.clicked.connect(self.scan_files_clicked)
        buttons.addWidget(self.scan_files)
        vertical.addLayout(buttons)
        return vertical

    def export_results_clicked(self):
        target_path = QFileDialog.getSaveFileName(self, 'Export Results', self._last_dir(), 'Excel (*.csv)')[0]
        if not target_path:
            return
        self._remember_dir(target_path)
        try:
            with open(target_path, 'wt', newline='') as csv_file:
                csv_writer = writer(csv_file)
                csv_writer.writerow(self.file_headers)
                for row in range(self.files.rowCount()):
                    csv_writer.writerow(
                        [self.files.item(row, 0).text()]
                        + [self.files.item(row, i + 1).text() for i in range(self.keyword_list.count())]
                    )
        except OSError as exc:
            QMessageBox.critical(self, 'Export Failed', f'Could not write results:\n{exc}')

    def export_excel_clicked(self):
        target_path = QFileDialog.getSaveFileName(self, 'Export Results', self._last_dir(), 'Excel Workbook (*.xlsx)')[
            0
        ]
        if not target_path:
            return
        self._remember_dir(target_path)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Results'
        sheet.append(self.file_headers)
        for row in range(self.files.rowCount()):
            values = []
            fills = []
            for column in range(self.files.columnCount()):
                item = self.files.item(row, column)
                values.append(item.text() if item else '')
                fills.append(item.background().color() if item else None)
            sheet.append(values)
            excel_row = row + 2  # header occupies row 1
            for column, color in enumerate(fills, start=1):
                if color is not None and color.isValid() and color.alpha() > 0:
                    hex_color = color.name(QColor.NameFormat.HexRgb).lstrip('#').upper()
                    sheet.cell(row=excel_row, column=column).fill = PatternFill(
                        start_color=hex_color, end_color=hex_color, fill_type='solid'
                    )
        try:
            workbook.save(target_path)
        except OSError as exc:
            QMessageBox.critical(self, 'Export Failed', f'Could not write results:\n{exc}')

    def save_keywords_clicked(self):
        target_path = QFileDialog.getSaveFileName(self, 'Save Keywords', self._last_dir(), 'Text (*.txt)')[0]
        if not target_path:
            return
        self._remember_dir(target_path)
        keywords = [self.keyword_list.item(i).text() for i in range(self.keyword_list.count())]
        try:
            with open(target_path, 'wt') as keywords_file:
                keywords_file.write('\n'.join(keywords) + '\n')
        except OSError as exc:
            QMessageBox.critical(self, 'Save Failed', f'Could not save keywords:\n{exc}')

    def load_keywords_clicked(self):
        source_path = QFileDialog.getOpenFileName(self, 'Load Keywords', self._last_dir(), 'Text (*.txt)')[0]
        if not source_path:
            return
        self._remember_dir(source_path)
        try:
            with open(source_path, 'rt') as keywords_file:
                loaded_keywords = [line for line in keywords_file.read().split('\n') if line]
        except OSError as exc:
            QMessageBox.critical(self, 'Load Failed', f'Could not load keywords:\n{exc}')
            return
        self.keyword_list.addItems(loaded_keywords)
        self._refresh_keyword_validity_highlighting()
        self.update_file_headers()
        if self.keyword_list.count():
            self.save_keywords.setDisabled(False)
        self.update_scan_button_state()

    def _match_mode(self):
        return self.match_mode_combo.currentData()

    def scan_files_clicked(self):
        if not self.file_names:
            return
        keywords = [self.keyword_list.item(i).text().lower() for i in range(self.keyword_list.count())]
        mode = self._match_mode()
        if mode is MatchMode.REGEX:
            for keyword in keywords:
                try:
                    re.compile(keyword)
                except re.error as exc:
                    QMessageBox.critical(
                        self, 'Invalid Regex', f'Keyword "{keyword}" is not a valid regex pattern:\n{exc}'
                    )
                    return
        self.scan_errors = []
        self.scan_files.setDisabled(True)
        self.add_files.setDisabled(True)
        # Sorting stays off for the duration of the scan: re-sorting on every
        # single incoming result would visibly shuffle rows mid-scan and cost
        # O(n log n) per result on a large batch. Re-enabled once it's done.
        self.files.setSortingEnabled(False)
        self.progress_dialog = QProgressDialog('Scanning files...', 'Cancel', 0, len(self.file_names), self)
        self.progress_dialog.setWindowTitle('Scanning')
        self.progress_dialog.setMinimumDuration(0)
        self.progress_dialog.setValue(0)
        self.scan_worker = ScanWorker(keywords, list(self.file_names), mode, self)
        self.scan_worker.result_ready.connect(self.handle_scan_result)
        self.scan_worker.finished_scanning.connect(self.handle_scan_finished)
        self.progress_dialog.canceled.connect(self.scan_worker.cancel)
        self.scan_worker.start()

    def handle_scan_result(self, result: ScanResult):
        row = self.file_items[result.file].row()
        if result.error is not None:
            self.scan_errors.append((result.file, result.error))
            for column in range(1, self.files.columnCount()):
                widget = QTableWidgetItem(self.options.display.error_text)
                widget.setBackground(self.options.display.error_color)
                widget.setToolTip(result.error)
                self.files.setItem(row, column, widget)
        else:
            for i, has_match in enumerate(result.matches):
                widget = QTableWidgetItem(self.options.display.found_text if has_match else self.options.display.missing_text)
                widget.setBackground(self.options.display.found_color if has_match else self.options.display.missing_color)
                if has_match:
                    widget.setToolTip(_build_tooltip(result.occurrences[i], result.truncated[i]))
                self.files.setItem(row, i + 1, widget)
        self.progress_dialog.setValue(self.progress_dialog.value() + 1)

    def handle_scan_finished(self):
        self.progress_dialog.close()
        self.update_scan_button_state()
        self.add_files.setDisabled(False)
        self.export_results.setDisabled(False)
        self.export_excel.setDisabled(False)
        self.files.setSortingEnabled(True)
        self._apply_filter()
        if self.scan_errors:
            details = '\n'.join(f'{basename(file)}: {error}' for file, error in self.scan_errors)
            QMessageBox.warning(
                self,
                'Some files failed',
                f'{len(self.scan_errors)} file(s) could not be scanned:\n\n{details}',
            )

    def update_scan_button_state(self):
        regex_ok = self._match_mode() is not MatchMode.REGEX or self._keywords_are_valid_regex()
        self.scan_files.setDisabled(not (self.file_names and self.keyword_list.count() and regex_ok))

    @staticmethod
    def _is_valid_regex(pattern):
        try:
            re.compile(pattern)
        except re.error:
            return False
        return True

    def _keywords_are_valid_regex(self):
        return all(
            self._is_valid_regex(self.keyword_list.item(i).text()) for i in range(self.keyword_list.count())
        )

    def _update_keyword_input_validity(self):
        text = self.new_keyword_text.text()
        invalid = self._match_mode() is MatchMode.REGEX and text and not self._is_valid_regex(text)
        if invalid:
            background = self.options.display.invalid_regex_background.name()
            foreground = self.options.display.invalid_regex_text.name()
            # Set both colors explicitly - a background-only stylesheet left
            # the text color to the OS theme, unreadable against a fixed
            # light background in dark mode.
            self.new_keyword_text.setStyleSheet(f'background-color: {background}; color: {foreground};')
        else:
            self.new_keyword_text.setStyleSheet('')
        self.add_keyword.setDisabled(not text or invalid)

    def _refresh_keyword_validity_highlighting(self):
        regex_mode = self._match_mode() is MatchMode.REGEX
        for i in range(self.keyword_list.count()):
            item = self.keyword_list.item(i)
            invalid = regex_mode and not self._is_valid_regex(item.text())
            if invalid:
                item.setBackground(self.options.display.invalid_regex_background)
                item.setForeground(self.options.display.invalid_regex_text)
            else:
                item.setBackground(QBrush())
                item.setForeground(QBrush())

    def _on_match_mode_changed(self):
        self._refresh_keyword_validity_highlighting()
        self._update_keyword_input_validity()
        self.update_scan_button_state()

    def add_files_clicked(self):
        files = QFileDialog.getOpenFileNames(
            self,
            'Files to Scan',
            self._last_dir(),
            'PDF Documents (*.pdf);;Word Documents (*.docx);;Text (*.txt);;All Files (*)',
        )[0]
        if files:
            self._remember_dir(files[0])
        self._add_files(files)

    def _add_files(self, files):
        if not files:
            return
        # Disable sorting during the bulk insert - avoids an O(n log n) resort
        # per row and any reordering mid-populate; restore + let it resort once.
        was_sorting = self.files.isSortingEnabled()
        self.files.setSortingEnabled(False)
        count = self.files.rowCount()
        self.files.setRowCount(count + len(files))
        for i, file in enumerate(files):
            self.file_names.append(file)
            item = QTableWidgetItem(basename(file))
            item.setData(FILE_PATH_ROLE, file)
            self.files.setItem(i + count, 0, item)
            self.file_items[file] = item
        self.files.setSortingEnabled(was_sorting)
        self.update_scan_button_state()
        self._apply_filter()

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()

    def dropEvent(self, event):
        files = [
            url.toLocalFile()
            for url in event.mimeData().urls()
            if url.isLocalFile() and isfile(url.toLocalFile())
        ]
        self._add_files(files)

    def remove_files_clicked(self):
        rows = sorted({index.row() for index in self.files.selectedIndexes()}, reverse=True)
        for row in rows:
            file = self.files.item(row, 0).data(FILE_PATH_ROLE)
            self.file_names.remove(file)
            del self.file_items[file]
            self.files.removeRow(row)
        self.remove_files.setDisabled(True)
        self.update_scan_button_state()

    def _apply_filter(self, *_args):
        query = self.filter_text.text().strip().lower()
        for row in range(self.files.rowCount()):
            if not query:
                self.files.setRowHidden(row, False)
                continue
            matched = any(
                self.files.item(row, column) is not None and query in self.files.item(row, column).text().lower()
                for column in range(self.files.columnCount())
            )
            self.files.setRowHidden(row, not matched)

    def _open_file_at_row(self, row, _column):
        item = self.files.item(row, 0)
        file = item.data(FILE_PATH_ROLE) if item else None
        if not file or not isfile(file):
            QMessageBox.warning(self, 'File Not Found', f'{file or "This file"} no longer exists.')
            return
        QDesktopServices.openUrl(QUrl.fromLocalFile(file))

    def create_keyword_area(self):
        horizontal = QHBoxLayout()
        self.new_keyword_text = QLineEdit()
        self.new_keyword_text.setPlaceholderText('Key word / phrase')
        self.new_keyword_text.textEdited.connect(self._update_keyword_input_validity)
        horizontal.addWidget(self.new_keyword_text)
        self.add_keyword = QPushButton('Add')
        self.add_keyword.setDisabled(True)
        self.add_keyword.clicked.connect(self.add_keyword_clicked)
        horizontal.addWidget(self.add_keyword)
        self.new_keyword_text.returnPressed.connect(self.add_keyword.click)
        vertical = QVBoxLayout()
        vertical.addLayout(horizontal)
        mode_row = QHBoxLayout()
        mode_row.addWidget(QLabel('Match mode:'))
        self.match_mode_combo = QComboBox()
        self.match_mode_combo.addItem('Substring', MatchMode.SUBSTRING)
        self.match_mode_combo.addItem('Whole word', MatchMode.WHOLE_WORD)
        self.match_mode_combo.addItem('Regex', MatchMode.REGEX)
        self.match_mode_combo.currentIndexChanged.connect(self._on_match_mode_changed)
        mode_row.addWidget(self.match_mode_combo)
        vertical.addLayout(mode_row)
        self.keyword_list = QListWidget()
        # No in-place editing: an edited keyword wouldn't refresh file headers,
        # save-state, or regex-validity highlighting. Remove + re-add instead.
        self.keyword_list.setEditTriggers(QListWidget.EditTrigger.NoEditTriggers)
        self.keyword_list.itemSelectionChanged.connect(
            lambda: self.remove_keyword.setDisabled(not self.keyword_list.selectedIndexes())
        )
        vertical.addWidget(self.keyword_list)
        self.remove_keyword = QPushButton('Remove')
        self.remove_keyword.setDisabled(True)
        self.remove_keyword.clicked.connect(self.remove_keyword_clicked)
        vertical.addWidget(self.remove_keyword)
        return vertical

    def add_keyword_clicked(self):
        self.keyword_list.addItem(self.new_keyword_text.text())
        self.new_keyword_text.setText('')
        self._update_keyword_input_validity()
        self.update_file_headers()
        self.save_keywords.setDisabled(False)
        self.update_scan_button_state()
        self.new_keyword_text.setFocus()

    def update_file_headers(self):
        self.file_headers = ['File'] + [self.keyword_list.item(i).text() for i in range(self.keyword_list.count())]
        old_columns = self.files.columnCount()
        len_headers = len(self.file_headers)
        self.files.setColumnCount(len_headers)
        self.files.setHorizontalHeaderLabels(self.file_headers)
        if old_columns < len_headers:
            header = self.files.horizontalHeader()
            for i in range(old_columns, len_headers):
                header.setSectionResizeMode(i, QHeaderView.ResizeMode.ResizeToContents)

    def remove_keyword_clicked(self):
        for item in self.keyword_list.selectedIndexes()[::-1]:
            self.keyword_list.takeItem(item.row())
        self.remove_keyword.setDisabled(True)
        self.update_file_headers()
        if not self.keyword_list.count():
            self.save_keywords.setDisabled(True)
        self.update_scan_button_state()
